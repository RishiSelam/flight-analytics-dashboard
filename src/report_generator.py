"""
Build reports/airline_kpi_report.xlsx with OpenPyXL.

Sheets: Executive Summary, Airline Performance, Route Analysis,
Airport Analysis, Passenger Analysis, Revenue Analysis.
"""

from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.config import REPORT_PATH

BLUE = "2A78D6"
DARK = "1C5CAB"
INK = "0B0B0B"
MUTED = "52514E"
ZEBRA = "F4F7FB"

TITLE_FONT = Font(size=16, bold=True, color=INK)
SUBTITLE_FONT = Font(size=10, italic=True, color=MUTED)
SECTION_FONT = Font(size=12, bold=True, color=DARK)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=DARK)
ZEBRA_FILL = PatternFill("solid", fgColor=ZEBRA)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(bottom=THIN)

PCT, INT, USD, USD2, MIN = "0.0%", "#,##0", "$#,##0", "$#,##0.00", "0.0"

LABELS = {
    "avg_delay_minutes": "Avg Delay (min)",
    "distance_miles": "Distance (mi)",
    "avg_ticket_price": "Avg Ticket Price",
    "departure_hour": "Departure Hour",
    "on_time_gap": "On-Time Gap vs Network",
    "passenger_mom_change": "Passengers MoM Change",
}


def _label(col: str) -> str:
    return LABELS.get(col, col.replace("_", " ").title().replace("On Time", "On-Time"))


def _number_format(col: str) -> str | None:
    if col.endswith(("_rate", "_share", "_gap", "_change")):
        return PCT
    if col in ("avg_ticket_price", "avg_fare", "revenue_per_flight"):
        return USD2
    if "revenue" in col:
        return USD
    if "minutes" in col:
        return MIN
    if col in ("departure_hour",) or col.endswith("_rank"):
        return "0"
    return INT


def write_table(ws: Worksheet, df: pd.DataFrame, row: int, title: str, note: str | None = None, col: int = 1) -> int:
    """Write a titled, formatted table. Returns the first free row after it."""
    ws.cell(row, col, title).font = SECTION_FONT
    row += 1
    if note:
        ws.cell(row, col, note).font = SUBTITLE_FONT
        row += 1

    for j, name in enumerate(df.columns):
        cell = ws.cell(row, col + j, _label(name))
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30
    header_row = row

    for i, record in enumerate(df.itertuples(index=False), start=1):
        for j, (name, value) in enumerate(zip(df.columns, record)):
            cell = ws.cell(row + i, col + j, None if pd.isna(value) else value)
            cell.border = BORDER
            if i % 2 == 0:
                cell.fill = ZEBRA_FILL
            if isinstance(value, (int, float)):
                cell.number_format = _number_format(name)

    first, last = header_row + 1, header_row + len(df)
    for j, name in enumerate(df.columns):
        letter = get_column_letter(col + j)
        rng = f"{letter}{first}:{letter}{last}"
        if name == "on_time_rate":
            ws.conditional_formatting.add(rng, ColorScaleRule(
                start_type="min", start_color="F8B4B4", mid_type="percentile", mid_value=50,
                mid_color="FFFFFF", end_type="max", end_color="B7E1C1"))
        elif name in ("revenue", "passengers") and len(df) > 3:
            ws.conditional_formatting.add(rng, DataBarRule(
                start_type="min", end_type="max", color="9EC5F4"))
    return last + 3


def _autofit(ws: Worksheet, min_width: int = 10, max_width: int = 28) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None or cell.font.i or (cell.font.sz or 11) >= 12:
                continue
            length = len(f"{cell.value:,.2f}") if isinstance(cell.value, float) else len(str(cell.value))
            widths[cell.column] = max(widths.get(cell.column, 0), length + 2)
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(min_width, min(width, max_width))


def _sheet(wb: Workbook, name: str, subtitle: str) -> Worksheet:
    ws = wb.create_sheet(name)
    ws["A1"] = name
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = SUBTITLE_FONT
    ws.sheet_view.showGridLines = False
    return ws


def _bar_chart(ws, title, data_col, cats_col, first_row, last_row, anchor, y_title, fmt, horizontal=True):
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.title, chart.y_axis.title = title, y_title
    chart.legend = None
    chart.y_axis.numFmt = fmt
    chart.y_axis.majorGridlines = None
    chart.add_data(Reference(ws, min_col=data_col, min_row=first_row - 1, max_row=last_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=cats_col, min_row=first_row, max_row=last_row))
    chart.series[0].graphicalProperties.solidFill = BLUE
    chart.series[0].graphicalProperties.line.solidFill = BLUE
    if horizontal:
        chart.x_axis.scaling.orientation = "maxMin"  # keep table order top-to-bottom
    chart.height, chart.width = 8, 16
    ws.add_chart(chart, anchor)


def _executive_summary(wb: Workbook, r: dict, insights: list[str], cleaning: dict) -> None:
    ws = wb.active
    ws.title = "Executive Summary"
    ws.sheet_view.showGridLines = False
    kpi = r["executive_kpis"].iloc[0]

    ws["A1"] = "Airline Operations & Passenger Analytics: KPI Report"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"Period {kpi.period_start} to {kpi.period_end}  |  "
                f"Generated {datetime.now():%Y-%m-%d %H:%M} by run_pipeline.py")
    ws["A2"].font = SUBTITLE_FONT

    ws["A4"] = "Headline KPIs"
    ws["A4"].font = SECTION_FONT
    kpis = [
        ("Total flights", kpi.total_flights, INT),
        ("Passengers", kpi.total_passengers, INT),
        ("Revenue", kpi.total_revenue, USD),
        ("Average fare", kpi.avg_fare, USD2),
        ("On-time rate (departure < 15 min late)", kpi.on_time_rate, PCT),
        ("Average departure delay (min)", kpi.avg_delay_minutes, MIN),
        ("Cancellation rate", kpi.cancellation_rate, PCT),
        ("Cancelled flights", kpi.cancelled_flights, INT),
    ]
    for cell, text in ((ws["A5"], "Metric"), (ws["B5"], "Value")):
        cell.value, cell.font, cell.fill = text, HEADER_FONT, HEADER_FILL
    for i, (label, value, fmt) in enumerate(kpis, start=6):
        ws.cell(i, 1, label).border = BORDER
        c = ws.cell(i, 2, float(value))
        c.number_format, c.border, c.font = fmt, BORDER, Font(bold=True)

    row = 6 + len(kpis) + 2
    ws.cell(row, 1, "Key Insights").font = SECTION_FONT
    for i, text in enumerate(insights, start=1):
        row += 1
        ws.cell(row, 1, f"{i}.").alignment = Alignment(vertical="top")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        cell = ws.cell(row, 2, text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 15 * (len(text) // 95 + 1) + 4

    row += 3
    ws.cell(row, 1, "Data Quality (cleaning log)").font = SECTION_FONT
    for key, value in cleaning.items():
        row += 1
        ws.cell(row, 1, key.replace("_", " ").capitalize()).border = BORDER
        c = ws.cell(row, 2, value)
        c.number_format, c.border = INT, BORDER

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 18
    for letter in "CDEFGH":
        ws.column_dimensions[letter].width = 14


def build_report(results: dict[str, pd.DataFrame], insights: list[str], cleaning: dict) -> None:
    r = results
    wb = Workbook()
    _executive_summary(wb, r, insights, cleaning)

    # Airline Performance
    ws = _sheet(wb, "Airline Performance", "On-time rate and delays measured over operated flights.")
    airlines = r["airline_performance"]
    end = write_table(ws, airlines, 4, "Airline scorecard")
    _bar_chart(ws, "On-time rate by airline", 3, 1, 6, 5 + len(airlines), "M4", "On-time rate", "0%")
    end = write_table(ws, r["delay_by_time_of_day"], end, "Performance by time of day")
    write_table(ws, r["monthly_trends"], end, "Monthly operational trends")

    # Route Analysis
    ws = _sheet(wb, "Route Analysis", "Directional routes (ORIGIN-DESTINATION), all carriers combined.")
    end = write_table(ws, r["high_demand_poor_performance"], 4,
                      "High demand, poor performance",
                      "Top passenger quartile routes with on-time rate below the network average.")
    end = write_table(ws, r["top_delayed_routes"], end, "Top 10 delayed routes", "Routes with 100+ flights.")
    end = write_table(ws, r["route_performance"], end, "All routes, ranked by revenue")
    write_table(ws, r["route_airline_breakdown"], end, "Shared routes: carrier comparison")

    # Airport Analysis
    ws = _sheet(wb, "Airport Analysis",
                "Departure metrics attributed to origin. Disruption = cancelled or 60+ min late.")
    airports = r["airport_performance"]
    end = write_table(ws, airports, 4, "Airport disruption ranking")
    _bar_chart(ws, "Disruption rate by airport", 6, 1, 6, 5 + len(airports), "O4", "Disruption rate", "0%")
    end = write_table(ws, r["airport_seasonality"], end, "Seasonal on-time rate by airport")
    write_table(ws, r["delay_by_hour"], end, "Delay by scheduled departure hour")

    # Passenger Analysis
    ws = _sheet(wb, "Passenger Analysis", "Passengers on operated flights, by month and cabin.")
    monthly = r["monthly_passengers"]
    end = write_table(ws, monthly, 4, "Monthly passengers")
    chart = LineChart()
    chart.title, chart.y_axis.title, chart.legend = "Monthly passengers", "Passengers", None
    chart.y_axis.numFmt = "#,##0"
    chart.add_data(Reference(ws, min_col=2, min_row=5, max_row=5 + len(monthly)), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=6, max_row=5 + len(monthly)))
    chart.series[0].graphicalProperties.line.solidFill = BLUE
    chart.series[0].graphicalProperties.line.width = 25000
    chart.height, chart.width = 8, 16
    ws.add_chart(chart, "K4")
    write_table(ws, r["cabin_class_summary"], end, "Cabin class distribution")

    # Revenue Analysis
    ws = _sheet(wb, "Revenue Analysis", "Revenue = passengers x ticket price.")
    airline_rev = r["airline_revenue"]
    end = write_table(ws, airline_rev, 4, "Revenue by airline")
    _bar_chart(ws, "Revenue by airline", 3, 1, 6, 5 + len(airline_rev), "I4", "Revenue", "$#,##0,,\"M\"")
    end = write_table(ws, r["top_revenue_routes"], end, "Top 10 revenue routes")
    cabin = r["cabin_class_summary"][["cabin_class", "revenue", "revenue_share", "avg_ticket_price"]]
    write_table(ws, cabin, end, "Revenue by cabin class")

    for sheet in wb.worksheets[1:]:
        _autofit(sheet)
        sheet.freeze_panes = "A4"

    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(REPORT_PATH)
